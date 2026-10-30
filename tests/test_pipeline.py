"""Unit and end-to-end tests. Run: python -m pytest -q"""
from __future__ import annotations

import subprocess
import sys
from datetime import datetime
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from csp_gps.extract import extract, is_relevant, mask_phone  # noqa: E402
from csp_gps.parser import parse_export_text, timestamp_from_filename  # noqa: E402


# --- parser ----------------------------------------------------------------------------------------
def test_android_12h_with_narrow_space_and_attachment():
    text = ("22/09/2026, 7:32 pm - Ravi: IMG-20260922-WA0012.jpg (file attached)\n"
            "location 827 m away\n"
            "22/09/2026, 7:40 pm - Ravi: second\n")
    msgs, unparsed = parse_export_text(text, "G", "g.txt")
    assert unparsed == 0 and len(msgs) == 2
    assert msgs[0].timestamp == datetime(2026, 9, 22, 19, 32)
    assert msgs[0].attachments == ["IMG-20260922-WA0012.jpg"]
    assert msgs[0].text == "location 827 m away"


def test_ios_bracket_format_and_attached_tag():
    text = "[02/10/26, 14:20:05] Neha: ‎<attached: 00000007-PHOTO-2026-10-02-14-20-05.jpg>\n"
    msgs, _ = parse_export_text(text, "G", "g.txt")
    assert msgs[0].timestamp == datetime(2026, 10, 2, 14, 20, 5)
    assert msgs[0].attachments == ["00000007-PHOTO-2026-10-02-14-20-05.jpg"]


def test_whatsapp_web_copy_format():
    msgs, _ = parse_export_text("[3:20 pm, 13/10/2026] Meena: GPS weak\n", "G", "g.txt")
    assert msgs[0].timestamp == datetime(2026, 10, 13, 15, 20) and msgs[0].sender == "Meena"


def test_month_first_detected_from_file():
    text = "9/22/26, 10:00 - A: x\n9/23/26, 10:00 - A: y\n"
    msgs, _ = parse_export_text(text, "G", "g.txt")
    assert msgs[0].timestamp.month == 9 and msgs[0].timestamp.day == 22


def test_system_message_has_no_sender_and_media_omitted():
    text = ("22/09/2026, 9:00 am - Messages and calls are end-to-end encrypted.\n"
            "22/09/2026, 9:01 am - A: <Media omitted>\n")
    msgs, _ = parse_export_text(text, "G", "g.txt")
    assert msgs[0].is_system and msgs[1].media_omitted


def test_timestamp_from_image_filename():
    assert timestamp_from_filename("WhatsApp Image 2026-09-22 at 7.32.26 PM.jpeg") == datetime(2026, 9, 22, 19, 32, 26)
    assert timestamp_from_filename("WhatsApp Image 2026-09-29 at 20.37.12 - Copy.jpeg") == datetime(2026, 9, 29, 20, 37, 12)
    assert timestamp_from_filename("holiday.jpg") is None


# --- extraction ---------------------------------------------------------------------------------------
def test_distance_csp_and_district():
    ex = extract("CSP ID: KO123456 location 827 m away, visit submit nahi ho raha. District: Patna", "")
    assert ex.relevant and ex.issue_type == "Distance Exceeds Radius"
    assert ex.distance_m == 827 and ex.csp_id == "KO123456" and ex.district == "Patna"
    assert ex.visit_status == "Not Completed / Not Updated"


def test_km_converted_and_distance_word_not_district():
    ex = extract("distance 1.2 km door bata raha hai", "")
    assert ex.distance_m == 1200 and ex.district == "Not Available"


def test_labelled_coordinates():
    ex = extract("Current location: 25.6093, 85.1376 Planned location: 25.6150, 85.1410", "")
    assert ex.actual_location == "25.6093, 85.1376" and ex.planned_location == "25.6150, 85.1410"


def test_hindi_gps_not_detected():
    ex = extract("लोकेशन नहीं आ रहा है, जिला: Muzaffarpur", "")
    assert ex.relevant and ex.issue_type == "GPS Not Detected" and ex.district == "Muzaffarpur"


def test_missing_values_are_not_invented():
    ex = extract("GPS location not detected", "")
    assert ex.csp_id == ex.district == ex.csp_name == "Not Available" and ex.distance_m is None


def test_not_relevant_messages():
    assert not is_relevant("Good morning team")
    assert not is_relevant("Please share CSP ID and district with every location issue.")
    assert not is_relevant("CSP ID: X1 location issue resolved, visit ho gaya")
    assert not is_relevant("Server error on payroll portal")      # error unrelated to a CSP visit
    assert is_relevant("Server error aa raha hai visit submit karte time")


def test_phone_masking():
    assert mask_phone("+91 98765 43210") == "+91 98XXX XX210"
    assert "43210" not in mask_phone("call 9876543210 now")


# --- end to end on synthetic data -------------------------------------------------------------------
@pytest.fixture(scope="module")
def synthetic_run(tmp_path_factory):
    out = tmp_path_factory.mktemp("out")
    state = ROOT / "data" / "state_synthetic.sqlite"
    state.unlink(missing_ok=True)
    py = sys.executable
    subprocess.run([py, str(ROOT / "scripts" / "generate_synthetic_data.py")], check=True, cwd=ROOT)
    subprocess.run([py, str(ROOT / "process.py"), "--synthetic", "--output", str(out)], check=True, cwd=ROOT)
    return out


def test_end_to_end_outputs(synthetic_run):
    df = pd.read_csv(synthetic_run / "issues.csv", dtype=str, keep_default_na=False)
    assert len(df) == 16
    assert set(df["Data Source"]) == {"SYNTHETIC TEST DATA"}
    assert all(df["Record ID"].str.startswith("SYN-"))
    assert "2026-09-20" not in set(df["Message Date"])                 # before the reporting period
    assert not df["Reported By"].str.contains("43210").any()            # phone masked
    dup = df[df["Duplicate Flag"] == "Possible Duplicate"]
    assert len(dup) == 1 and "same image file" in dup.iloc[0]["Duplicate Reason"]
    statuses = dict(zip(df["CSP ID"] + df["Message Date"], df["Resolution Status"], strict=True))
    assert statuses["SYN100012026-09-22"] == "Resolved"
    assert statuses["SYN200012026-10-02"] == "Escalated"
    multi = df[df["CSP ID"] == "SYN10003"].iloc[0]
    assert multi["Evidence Filename"].count(";") == 1                   # caption + 2 photos = one record
    assert df.loc[df["CSP ID"] == "SYN10007", "Verification Reasons"].str.contains("Media missing").all()
    assert (df["Distance Mismatch (metres)"] == "").sum() == 11          # missing stays blank, not 0


def test_workbook_structure(synthetic_run):
    wb = load_workbook(synthetic_run / "SYNTHETIC_TEST_CSP_GPS_Issue_Report.xlsx")
    assert wb.sheetnames[0] == "Daily Update"
    assert wb.sheetnames[1:8] == ["Executive Summary", "Detailed Issue Tracker", "Issue Category Analysis",
                                 "District and CSP Analysis", "Resolution Tracker", "Data Quality and Verification",
                                 "Evidence Index"]
    ws = wb["Detailed Issue Tracker"]
    assert ws.freeze_panes == "C4" and ws.auto_filter.ref
    assert [c.value for c in ws[3]][:3] == ["Record ID", "Message Date", "Message Time"]
    assert len(ws.data_validations.dataValidation) == 3
    assert "SYNTHETIC" in ws["A2"].value


def test_rerun_keeps_ids_and_manual_edits(synthetic_run):
    report = synthetic_run / "SYNTHETIC_TEST_CSP_GPS_Issue_Report.xlsx"
    wb = load_workbook(report)
    ws = wb["Detailed Issue Tracker"]
    header = [c.value for c in ws[3]]
    ws.cell(row=5, column=header.index("Resolution Status") + 1, value="In Progress")
    ws.cell(row=5, column=header.index("Assigned To") + 1, value="Test Owner")
    rid = ws.cell(row=5, column=1).value
    wb.save(report)
    before = pd.read_csv(synthetic_run / "issues.csv", dtype=str)["Record ID"].tolist()
    subprocess.run([sys.executable, str(ROOT / "process.py"), "--synthetic", "--output", str(synthetic_run)],
                   check=True, cwd=ROOT)
    df = pd.read_csv(synthetic_run / "issues.csv", dtype=str, keep_default_na=False)
    assert df["Record ID"].tolist() == before                           # no new/duplicate records
    row = df[df["Record ID"] == rid].iloc[0]
    assert row["Resolution Status"] == "In Progress" and row["Assigned To"] == "Test Owner"


# --- patterns seen in the real "Kiosk Team Eko" export -------------------------------------------
def test_unlabelled_eko_csp_code_and_gps_issue():
    ex = extract("@~Amit singh 1A851811  GPS Issue", "")
    assert ex.relevant and ex.csp_id == "1A851811" and ex.issue_type == "GPS Not Detected"


def test_checkins_notices_and_form_labels_are_not_complaints():
    assert not is_relevant("Leaving location for CSP visit Fatehabad")
    assert not is_relevant("Live location CSP visit Katihar RBO")
    assert not is_relevant("Csp code. 1A851604\nLocation. Siswa awwal\n@Ganesh add account")
    assert not is_relevant("Ganesh sir will coordinate")
    assert not is_relevant("Regarding the location issue, please ensure exact latitude is updated without fail")


def test_location_update_and_misspelt_distance():
    assert extract("1A850396 Please Update Location ditels", "").issue_type == "CSP Location Needs Update"
    ex = extract("Code -1A852249\nDistance - 50 Miter\nDistance -500 Miter", "")
    assert ex.distance_m == 50 and ex.issue_type == "Distance Exceeds Radius"


def test_markdown_reply_quotes_link_to_messages(tmp_path):
    from csp_gps.parser import attach_markdown_quotes
    msgs, _ = parse_export_text("[10/13/26, 9:46:41 AM] Sakshi: @~Amit 1A851811 GPS Issue\n"
                                "[10/13/26, 10:50:12 AM] Amit: Issue resolved\n", "G", "g.txt")
    md = tmp_path / "chat.md"
    md.write_text("# WhatsApp Chat Export: Kiosk Team\n\n## October 13, 2026\n\n"
                  "[9:46 AM] **Sakshi:** @~Amit 1A851811 GPS Issue\n\n"
                  "[10:50 AM] **Amit:**\n> _Sakshi: @~Amit 1A851811 GPS Issue_\nIssue resolved\n", encoding="utf-8")
    title, linked = attach_markdown_quotes(msgs, md)
    assert title == "Kiosk Team" and linked == 1 and "1A851811" in msgs[1].quoted


# --- patterns seen in the 20 Oct exports (DC one-to-one chat, daily reports) --------------------
def test_location_error_and_offline_daily_report_items():
    assert extract("1A852410 anup sanchla visit location error", "").issue_type == "Location Error in Visit App"
    ex = extract("2. Savita kumari -1A850534( Offline)", "")
    assert ex.relevant and ex.csp_id == "1A850534" and ex.issue_type == "Visit Not Completed/Updated"
    assert extract("1A851943\nremapping ka issue resolve karwiaye", "").issue_type == "CSP Not in Visit App"


def test_forwarded_tag_is_ignored_when_comparing_text():
    from csp_gps.pipeline import _norm_text
    assert _norm_text("[Forwarded] 1A851643 prem offline visit done") == _norm_text("1A851643 prem offline visit done")


# --- agent brief ----------------------------------------------------------------------------------
def test_agent_brief_drafts_and_changes():
    from datetime import date as d

    from csp_gps.agent_brief import build_brief
    rows = [
        # id, csp, type, status, step, by, days, evidence
        ("GPS-1", "1A850001", "GPS Not Detected", "Open", "Escalate - pending 9 days", "DC A", 9, ""),
        ("GPS-2", "1A850001", "App/Server Error", "Open", "Escalate - pending 8 days", "DC A", 8, ""),
        ("GPS-3", "1A850002", "CSP Not in Visit App", "Open", "Follow up - pending 3 days", "DC B", 3, ""),
        ("GPS-4", "Not Available", "Visit Not Completed/Updated", "Open", "New - take suggested action", "DC C", 1, ""),
        ("GPS-5", "1A850003", "GPS Not Detected", "Resolved", "Done", "DC A", None,
         "Auto-closed by agent: SBI CSP Visit App shows a visit"),
    ]
    issues = pd.DataFrame([{"Record ID": r[0], "CSP ID": r[1], "Issue Type": r[2], "Resolution Status": r[3],
                            "Next Step": r[4], "Reported By": r[5], "Days Pending": r[6], "Resolution Evidence": r[7],
                            "Duplicate Flag": "Unique", "Message Date": d(2026, 10, 10), "Issue Description": "visit nahi hua"}
                           for r in rows])
    prev = pd.DataFrame({"Record ID": ["GPS-1", "GPS-2", "GPS-3", "GPS-5"],
                         "Resolution Status": ["Open", "Open", "Open", "Open"],
                         "Next Step": ["Escalate", "Follow up", "New", "New"]})
    b = build_brief(issues, prev, {"tech_support": "Amit"}, today=d(2026, 10, 21))
    c = b["counts"]
    assert c["new"] == 1 and c["auto_closed"] == 1 and c["closed_since_last_run"] == 1 and c["newly_escalate"] == 1
    support = next(x for x in b["drafts"] if x["to"].startswith("@Amit"))
    assert support["message"].count("1A850001") == 1          # one line per CSP
    assert any(x["to"] == "@DC C" for x in b["drafts"])         # asked for the missing CSP ID
    assert b["drafts"][0]["to"] == "Ma'am"


# --- Docker report scheduler -------------------------------------------------------------------------
def test_scheduler_next_run():
    from datetime import datetime
    from zoneinfo import ZoneInfo

    sys.path.insert(0, str(ROOT / "scripts"))
    from scheduler import next_run, parse_times

    times = parse_times("15:00, 10:00,10:00")
    assert times == [(10, 0), (15, 0)]
    ist = ZoneInfo("Asia/Kolkata")
    assert next_run(datetime(2026, 10, 21, 9, 59, tzinfo=ist), times) == datetime(2026, 10, 21, 10, 0, tzinfo=ist)
    assert next_run(datetime(2026, 10, 21, 10, 0, tzinfo=ist), times) == datetime(2026, 10, 21, 15, 0, tzinfo=ist)
    assert next_run(datetime(2026, 10, 21, 16, 30, tzinfo=ist), times) == datetime(2026, 10, 22, 10, 0, tzinfo=ist)
    with pytest.raises(ValueError):
        parse_times("25:00")


# --- SBI visit report: score 0 = not completed --------------------------------------------------------
def test_score_zero_visits_not_counted(tmp_path):
    from csp_gps.visits import done_visits, load_visit_reports, score_zero_visits

    folder = tmp_path / "visit_reports"
    folder.mkdir()
    pd.DataFrame({
        "CSP Code": ["1A000001", "1A000001", "1A000002", "1A000003", "1A000003"],
        "Visit Date": ["2026-10-08"] * 5,
        "Visit Score": [0, 90, 0, 85, 85],
        "username": ["BC1", "BC1", "BC2", "BC3", "BC3"],
    }).to_excel(folder / "Visit report test.xlsx", index=False)
    v, failed = load_visit_reports(folder, tmp_path)
    assert not failed
    assert (v["Duplicate Row"] == "No").sum() == 3            # each visit once
    done, zero = done_visits(v), score_zero_visits(v)
    assert sorted(done["CSP Code"]) == ["1A000001", "1A000003"]  # the 90 row is kept, not the 0 repeat
    assert done.loc[done["CSP Code"] == "1A000001", "Visit Score"].item() == 90
    assert list(zero["CSP Code"]) == ["1A000002"]


def test_personal_ids_masked_and_csp_forms_skipped():
    from csp_gps.extract import is_relevant, mask_phone

    out = mask_phone("Aadhaar 4159 5017 6835 PAN ABCDE1234F A/C 34028031763 mail ravi.k@gmail.com lat 26.6317 CSP 1A850933")
    assert "4159" not in out and "ABCDE1234F" not in out and "340280" not in out and "ravi.k@" not in out
    assert "6835" in out and "26.6317" in out and "1A850933" in out
    form = ("Code of Replacement National BC Code 1A850933 Gram Panchayat X Pin Code 274408 "
            "Aadhaar Number 1 PAN Number 2 Latitude 26.6317 Longitude 83.8982")
    assert not is_relevant(form)
    assert is_relevant("CSP visit 1A852434 Remark - CSP not mapped Latitude 27.724919 Longitude 80.440413")


# --- dashboard upload box --------------------------------------------------------------------------------
def test_upload_box_saves_checks_and_skips_repeats(tmp_path):
    import io
    import zipfile

    from csp_gps.uploads import save_upload, uploads_enabled

    assert uploads_enabled({"CSP_UPLOADS": "1"}) and not uploads_enabled({})
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("chat.txt", "[10/27/26, 4:00:00 PM] Ajay: 1A850933 location issue\n")
    export = buf.getvalue()
    first = save_upload("Kiosk Team Eko.zip", export, tmp_path)
    assert first.ok and (tmp_path / "Kiosk Team Eko.zip").exists()
    dup = save_upload("copy.zip", export, tmp_path)
    assert not dup.ok and dup.known                                       # same content -> skipped, still rebuilds
    assert not save_upload("photo.zip", b"not a zip", tmp_path).ok          # not an export
    assert not save_upload("../../evil.exe", b"x", tmp_path).ok             # wrong type
    assert save_upload("../x/../Kiosk Team Eko.zip", export + b"", tmp_path).ok is False

    report = io.BytesIO()
    pd.DataFrame({"CSP Code": ["1A000001"], "Visit Date": ["2026-10-08"], "Visit Score": [90]}).to_excel(report, index=False)
    r = save_upload("Visit report 1st Oct to 28th Oct.xlsx", report.getvalue(), tmp_path)
    assert r.ok and (tmp_path / "visit_reports" / "Visit report 1st Oct to 28th Oct.xlsx").exists()
    other = io.BytesIO()
    pd.DataFrame({"Name": ["x"]}).to_excel(other, index=False)
    assert not save_upload("random.xlsx", other.getvalue(), tmp_path).ok


def test_each_csp_counted_once_in_visit_totals(tmp_path):
    from csp_gps.visits import done_csps, load_visit_reports, score_zero_csps

    folder = tmp_path / "visit_reports"
    folder.mkdir()
    pd.DataFrame({
        "CSP Code": ["1A000001", "1A000001", "1A000002", "1A000003", "1A000003"],
        "Visit Date": ["2026-10-08", "2026-10-20", "2026-10-09", "2026-10-10", "2026-10-21"],
        "Visit Score": [88, 90, 0, 0, 85],
        "username": ["BC1", "BC1", "BC2", "BC3", "BC3"],
    }).to_excel(folder / "Visit report.xlsx", index=False)
    v, _ = load_visit_reports(folder, tmp_path)
    assert sorted(done_csps(v)["CSP Code"]) == ["1A000001", "1A000003"]   # visited twice -> counted once
    assert list(score_zero_csps(v)["CSP Code"]) == ["1A000002"]            # 1A000003 later completed


# --- CSP location list ------------------------------------------------------------------------------------
def test_csp_locations_load_mask_and_merge(tmp_path):
    from csp_gps.locations import load_locations, merge_locations
    from csp_gps.uploads import save_upload

    csv = (b"CSP Code,Name,Number,\"Lat, Long\",Other reply\n"
           b"1A852327,Anuj Kumar,917983534146,\"28.9486537, 78.5371186\",\n"
           b"1A000009,Far Away,919000000001,\"51.5, -0.12\",sent an image\n")
    r = save_upload("locations.csv", csv, tmp_path / "input")
    assert r.ok and (tmp_path / "input" / "csp_locations" / "locations.csv").exists()
    loc = load_locations(tmp_path / "input" / "csp_locations", tmp_path)
    assert len(loc) == 2 and "7983534146" not in " ".join(loc["Number"])
    assert dict(zip(loc["CSP Code"], loc["Location Check"], strict=True))["1A000009"].startswith("Check")
    issues = pd.DataFrame({"CSP ID": ["1A852327"], "Duplicate Flag": ["Unique"], "Resolution Status": ["Open"]})
    done = pd.DataFrame({"CSP Code": ["1A852327"], "Visit Date": ["2026-10-24"]})
    m = merge_locations(loc, issues, done).set_index("CSP Code")
    assert m.loc["1A852327", "Pending Complaints"] == 1 and m.loc["1A852327", "Last App Visit"] == "2026-10-24"
    assert m.loc["1A852327", "Map"].startswith("https://maps.google.com/?q=28.9486537")
